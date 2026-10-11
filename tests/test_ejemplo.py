"""El ejemplo: los tres medios salen de los tokens y el tablero pasa el auditor."""
import re
from pathlib import Path

from openpyxl import load_workbook

from adaptadores.html import auditar_html as A
from ejemplo import construir as E
from nucleo import tokens as T

RAIZ = Path(__file__).resolve().parents[1]


def test_tablero_conforme(tmp_path):
    ruta = E.tablero(tmp_path / "tablero.html")
    A.revisar(ruta)
    assert A.hallazgos == []


def test_tablero_usa_la_convencion_de_cifras(tmp_path):
    html = E.tablero(tmp_path / "t.html").read_text(encoding="utf-8")
    assert "103,4 %" in html                       # coma decimal y espacio antes del %
    total = sum(r[2] for r in E.REGIONES)
    assert f'class="cifra">{E.entero(total)}<' in html
    assert E.entero(6700) == "6.700"


def test_excel_resumen_primero_y_semaforo(tmp_path):
    ruta = E.excel(tmp_path / "s.xlsx")
    wb = load_workbook(ruta)
    assert wb.sheetnames[0] == "Resumen"
    ws = wb["Resumen"]
    assert ws["A1"].fill.fgColor.rgb.endswith(T.color("principal_oscuro").lstrip("#"))
    assert len(ws.conditional_formatting) >= 1


def test_grafica_png(tmp_path):
    ruta = E.grafica(tmp_path / "g.png")
    assert ruta.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


def test_construir_completo_devuelve_0():
    assert E.main() == 0


def test_la_paleta_de_la_guia_de_tableros_sale_de_los_tokens():
    # La guía no copia colores (se quedarían viejos al cambiar la marca): pide
    # roles a T.css_variables. Se comprueba que esos roles existen y que el bloque
    # generado lleva exactamente los colores de los tokens.
    doc = (RAIZ / "docs" / "tableros.md").read_text(encoding="utf-8")
    assert not re.findall(r"#[0-9A-Fa-f]{6}\b", doc)
    llamada = re.search(r"T\.css_variables\(\[(.*?)\]\)", doc, re.S).group(1)
    roles = re.findall(r'"(\w+)"', llamada)
    assert roles and set(roles) <= set(T.tokens()["roles"])
    css = T.css_variables(roles)
    for rol in roles:
        assert f"--{rol.replace('_', '-')}:{T.color(rol)};" in css


def test_css_variables_trae_todos_los_roles():
    css = T.css_variables()
    assert css.startswith(":root{") and css.endswith("}")
    assert css.count("--") == len(T.tokens()["roles"])
