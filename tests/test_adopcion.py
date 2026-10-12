"""Adoptar el sistema con otra marca: se cambia SOLO nucleo/tokens.json y todo
sigue funcionando. Se hace en una copia del repositorio, con una paleta que no
se parece a la de demostración."""
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]

# Otra marca (violeta), con los mismos roles. Cumple los contrastes de test_tokens.
OTRA_PALETA = {
    "TINTA": "#1E1B2E", "TINTA_SUAVE": "#4A4560", "AZUL": "#5B3E9E",
    "AZUL_CLARO": "#B9A6E3", "AZUL_PROFUNDO": "#2E1A63", "AMBAR": "#D9A21B",
    "AMBAR_CLARO": "#F3E1A9", "VERDE": "#2C7A4B", "VERDE_CLARO": "#C4E6D1",
    "ROJO": "#A8322D", "ROJO_CLARO": "#F2CBC8", "GRIS": "#76717F",
    "GRIS_CLARO": "#E7E5EC", "PAPEL": "#FAF9FC",
}
OTRAS_SERIES = {
    "secuencial": ["#E3DAF5", "#B9A6E3", "#7C5FC2", "#2E1A63"],
    "categorica": ["#5B3E9E", "#D9A21B", "#2C7A4B", "#A8322D", "#B9A6E3", "#76717F"],
    "semaforo": {"bien": "#2C7A4B", "alerta": "#D9A21B", "mal": "#A8322D"},
}

pytestmark = pytest.mark.skipif(os.environ.get("DENTRO_DE_ADOPCION") == "1",
                                reason="ya corre dentro de la copia")


@pytest.fixture(scope="module")
def copia(tmp_path_factory) -> Path:
    destino = tmp_path_factory.mktemp("otra_marca") / "repo"
    shutil.copytree(RAIZ, destino, ignore=shutil.ignore_patterns(
        ".git", ".venv", "venv", "salida", "__pycache__", ".pytest_cache", ".ruff_cache"))
    ruta = destino / "nucleo" / "tokens.json"
    t = json.loads(ruta.read_text(encoding="utf-8"))
    viejos = {v.upper() for v in t["paleta"].values()}
    t["paleta"], t["series"] = OTRA_PALETA, OTRAS_SERIES
    ruta.write_text(json.dumps(t, ensure_ascii=False, indent=2), encoding="utf-8")
    destino.joinpath("viejos.txt").write_text("\n".join(sorted(viejos)), encoding="utf-8")
    return destino


def _correr(copia: Path, *args) -> subprocess.CompletedProcess:
    # El modo --navegador no depende de la paleta y ya se prueba fuera: dentro de
    # la copia se omite para no abrir Chrome dos veces.
    entorno = {**os.environ, "DENTRO_DE_ADOPCION": "1", "AUDITOR_NAVEGADOR": "/no/existe"}
    return subprocess.run([sys.executable, *args], cwd=copia, env=entorno, capture_output=True,
                          text=True, encoding="utf-8", check=False, timeout=600)


def test_construir_pasa_con_otra_marca(copia):
    r = _correr(copia, "ejemplo/construir.py")
    assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
    assert "tablero conforme         → pasa" in r.stdout
    assert "tablero con defectos     → detectado" in r.stdout


def test_los_tres_medios_usan_la_marca_nueva(copia):
    _correr(copia, "ejemplo/construir.py")
    html = (copia / "ejemplo" / "salida" / "tablero.html").read_text(encoding="utf-8").upper()
    assert OTRA_PALETA["AZUL"] in html and OTRA_PALETA["PAPEL"] in html
    viejos = set((copia / "viejos.txt").read_text(encoding="utf-8").split())
    assert not {v for v in viejos if v in html} - set(OTRA_PALETA.values())
    ws = load_workbook(copia / "ejemplo" / "salida" / "seguimiento.xlsx")["Resumen"]
    assert ws["A1"].fill.fgColor.rgb.endswith(OTRA_PALETA["AZUL_PROFUNDO"].lstrip("#"))
    assert (copia / "ejemplo" / "salida" / "grafica.png").stat().st_size > 10_000


def test_todas_las_pruebas_pasan_con_otra_marca(copia):
    r = _correr(copia, "-m", "pytest", "-q", "-p", "no:cacheprovider")
    assert r.returncode == 0, r.stdout[-3000:]
