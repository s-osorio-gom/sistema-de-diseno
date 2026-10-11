#!/usr/bin/env python3
"""
excel.py — openpyxl con los tokens de la casa.

    from adaptadores.excel import excel as X
    ws = wb.active
    X.encabezado(ws, "A1:E1")
    X.zebra(ws, 2, 40, 5)
    X.semaforo(ws, "D2:D40", bien=95, mal=80)

Un Excel también se lee con los ojos y también cuenta una historia. Dos reglas
que se saltan siempre: **la hoja de resumen va primera**, no última —es el
top-down del Excel—, y el verde y el rojo salen de la paleta, no de los que
Excel trae por defecto.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from nucleo import tokens as T  # noqa: E402


def _hex(rol: str) -> str:
    return T.color(rol).lstrip("#").upper()


def fuente(nivel: str = "cuerpo", color: str = "texto", blanco: bool = False) -> Font:
    t = T.tokens()["tipografia"]
    return Font(name=t["familia"], size=T.pt(nivel) * 0.75,   # pt de lámina -> pt de celda
                bold=T.peso(nivel) >= 600,
                color="FFFFFF" if blanco else _hex(color))


def encabezado(ws, rango: str):
    """Fila de encabezado: fondo principal, texto blanco, congelada."""
    relleno = PatternFill("solid", fgColor=_hex("principal_oscuro"))
    for fila in ws[rango]:
        for c in fila:
            c.fill = relleno
            c.font = fuente("apoyo", blanco=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws[rango.split(":")[0]].offset(row=1).coordinate


def zebra(ws, fila_ini: int, fila_fin: int, n_col: int):
    """Bandas alternas. Sin ellas, el ojo salta de fila en una tabla ancha."""
    relleno = PatternFill("solid", fgColor=_hex("fondo"))
    for r in range(fila_ini, fila_fin + 1):
        if r % 2 == 0:
            for c in range(1, n_col + 1):
                ws.cell(row=r, column=c).fill = relleno


def semaforo(ws, rango: str, bien: float, mal: float):
    """Pinta la celda, no un icono al lado: el color tiene que estar en el dato."""
    s = T.tokens()["series"]["semaforo"]
    # El texto negro por defecto no alcanza 4,5:1 sobre el verde ni el rojo, y el
    # blanco no se lee sobre el ámbar: cada estado lleva su color de letra.
    blanco, tinta = Font(color="FFFFFF"), Font(color=_hex("texto"))
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="greaterThanOrEqual", formula=[str(bien)], font=blanco,
        fill=PatternFill("solid", fgColor=s["bien"].lstrip("#"))))
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="lessThanOrEqual", formula=[str(mal)], font=blanco,
        fill=PatternFill("solid", fgColor=s["mal"].lstrip("#"))))
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="between", formula=[str(mal), str(bien)], font=tinta,
        fill=PatternFill("solid", fgColor=s["alerta"].lstrip("#"))))


def destacar(celda):
    """La celda que importa, en color de acento. Una por hoja, no diez."""
    celda.fill = PatternFill("solid", fgColor=_hex("acento"))
    celda.font = fuente("destacado")


def bordes_suaves(ws, rango: str):
    lado = Side(style="thin", color=_hex("borde"))
    for fila in ws[rango]:
        for c in fila:
            c.border = Border(left=lado, right=lado, top=lado, bottom=lado)


def ancho(ws, anchos: dict[int, float]):
    for col, w in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = w
